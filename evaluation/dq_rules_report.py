"""Score the Data Quality rules stage against the answer key.

Usage, from the repo root:  python evaluation/dq_rules_report.py
Runs on a throwaway in-memory database, so backend/ams.db is never touched.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "backend"))

from app.core.config import settings  # noqa: E402
from app.db.session import Base  # noqa: E402
from app.models.ticket import AgentRun  # noqa: E402
from app.services.data_quality import run_data_quality  # noqa: E402
from app.services.ingestion import ingest_workbook  # noqa: E402


def answer_key() -> dict[str, dict]:
    rows = load_workbook(REPO / "evaluation" / "ground_truth.xlsx", read_only=True).active.iter_rows(values_only=True)
    header = next(rows)
    return {r[0]: dict(zip(header, r)) for r in rows}


def main() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    key = answer_key()

    with Session(engine) as db:
        ingest_workbook(db, settings.TICKETS_FILE)
        run_data_quality(db)
        runs = db.query(AgentRun).all()

    should_fail = [r for r in runs if key[r.ticket_id]["dq_should_pass"] == "No"]
    should_pass = [r for r in runs if key[r.ticket_id]["dq_should_pass"] == "Yes"]
    caught = [r for r in should_fail if r.verdict != "uncertain"]
    false_fails = [r for r in should_pass if r.verdict != "uncertain"]
    with_app = [r for r in should_pass if key[r.ticket_id]["expected_application"]]
    app_ok = [r for r in with_app if r.output["signals"]["app_found"] == key[r.ticket_id]["expected_application"]]
    shortest_pass = min(should_pass, key=lambda r: r.output["signals"]["body_length"])

    print(f"Open tickets checked      {len(runs)}")
    print(f"Incomplete caught         {len(caught)} / {len(should_fail)}")
    print(f"Good tickets wrongly held {len(false_fails)} / {len(should_pass)}")
    print(f"Left for the model        {len(runs) - len(caught) - len(false_fails)}")
    print(f"Application found         {len(app_ok)} / {len(with_app)}")
    print(f"Shortest good description {shortest_pass.output['signals']['body_length']} characters ({shortest_pass.ticket_id})")

    for title, items in [
        ("Incomplete, left for the model", [r for r in should_fail if r.verdict == "uncertain"]),
        ("Good, wrongly held", false_fails),
        ("Wrong application", [r for r in with_app if r not in app_ok]),
    ]:
        if items:
            print(f"\n{title}:")
            for r in items:
                print(f"  {r.ticket_id:<11} {r.verdict:<16} app={r.output['signals']['app_found']} "
                      f"expected={key[r.ticket_id]['expected_application']}  {r.output['reason']}")


if __name__ == "__main__":
    main()
