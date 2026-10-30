"""Mask personal information in the ServiceNow tickets workbook and write a masked copy.

Usage, from the repo root:  python backend/scripts/mask_servicenow_tickets.py
Reads  data/tickets/servicenow_tickets.xlsx
Writes data/tickets/servicenow_tickets_masked.xlsx (same columns; caller and free text masked)

One Masker covers the whole file, so the same person has the same placeholder in every ticket.
"""

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "backend"))

from app.privacy.masking import Masker  # noqa: E402

SOURCE = REPO / "data" / "tickets" / "servicenow_tickets.xlsx"
TARGET = REPO / "data" / "tickets" / "servicenow_tickets_masked.xlsx"
MASKED_FIELDS = ("caller", "short_description", "description", "resolution_notes")


def main() -> None:
    sheet = load_workbook(SOURCE, read_only=True)["servicenow_tickets"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    tickets = [dict(zip(header, r)) for r in rows if any(r)]

    masker = Masker(known_names=[t["caller"] for t in tickets])
    for t in tickets:
        for field in MASKED_FIELDS:
            if t[field] is not None:
                t[field] = str(masker.mask(t[field]))

    wb = Workbook()
    ws = wb.active
    ws.title = "servicenow_tickets"
    ws.append(header)
    for t in tickets:
        ws.append([t[h] for h in header])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(TARGET)

    print(f"Masked {len(tickets)} tickets -> {TARGET.relative_to(REPO)}")
    print(f"Findings: {dict(masker.counts.most_common())}")


if __name__ == "__main__":
    main()
